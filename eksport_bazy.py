# -*- coding: utf-8 -*-
"""Eksport bazy: jeden arkusz źródłowy (xlsx) -> data/data.json dla strony.

Użycie:
    python eksport_bazy.py "Księga radziecka Wiśnicza 1712-1736.xlsx"

Arkusz ma jeden wiersz nagłówka i bloki wpisów: wiersz z ID otwiera wpis,
kolejne wiersze bez ID dokładają osoby. Kolumny (nagłówki z arkusza):
ID, Dzień, Miesiąc, Rok, Strona, Skan, Kto, Z kim, Stan społeczny, Osoby powiązane,
Rola, Zawód, Skąd?, Uwagi, Wiek, Martwy?, Temat sprawy, Charakter wpisu,
Słowa klucze, Dodatkowe okoliczności, Położenie domu, Mięso?, Kobieta sprawcza?

Identyfikatory wpisów są stałe: wpis już obecny w data/data.json (ta sama data,
skan i kolejność w obrębie tej pary) zachowuje swój id, więc linki do wpisów
nie przestają działać. Zredagowane opisy okoliczności bierzemy z nota_red.json.
"""
import openpyxl, json, re, os, sys, unicodedata
from collections import defaultdict, Counter

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "Ksiega_radziecka_Wisnicza_1712-1736.xlsx")
OUT = os.path.join(HERE, "data", "data.json")
RED = os.path.join(HERE, "nota_red.json")

KOLUMNY = ["ID", "Dzień", "Miesiąc", "Rok", "Strona", "Skan", "Kto", "Z kim", "Stan społeczny",
           "Osoby powiązane", "Rola", "Zawód", "Skąd?", "Uwagi", "Wiek", "Martwy?", "Temat sprawy",
           "Charakter wpisu", "Słowa klucze", "Dodatkowe okoliczności", "Położenie domu",
           "Mięso?", "Kobieta sprawcza?"]

# ---------- słowa klucze: małe litery, bez kropek ----------
# Kropka ze spacją po słowie („czapka. dziwować się”) rozdziela dwa hasła; inicjały
# („P. J.”) nie rozdzielają. Hasła z KW_ORYGINAL zostają w pisowni źródła.
KW_SEP = re.compile(r"[,;/]+|(?<=\w\w)\.\s+")
KW_ORYGINAL = {"męka p j chrystusa": "męka P. J. Chrystusa"}
def norm_kw(t):
    n = re.sub(r"\s+", " ", (t or "").replace(".", " ")).strip().lower()
    return KW_ORYGINAL.get(n, n)
def split_kw(s):
    out = []
    for t in KW_SEP.split(s or ""):
        t = norm_kw(t)
        if t and t not in out:
            out.append(t)
    return out

# ---------- redakcja notatek (bez zmian względem export_json.py) ----------
def zdanie(s):
    s = (s or "").strip()
    if not s:
        return ""
    if s[-1] not in ".!?":
        s += "."
    return s[0].upper() + s[1:]

# żeńskie imiona najczęstsze w źródle (gdy nazwisko nie rozstrzyga)
_IMIONA_F = {"anna", "agnes", "agnieszka", "maria", "marianna", "regina", "zofia",
             "catharina", "katarzyna", "teresa", "teresua", "zuzanna", "susanna",
             "dorota", "ewa", "eva", "jadwiga", "hedwigis", "barbara", "elżbieta",
             "elisabeth", "magdalena", "helena", "krystyna", "christina", "konegunda",
             "conegunda", "kunegunda", "małgorzata", "margaretha"}
_IMIONA_M = {"jan", "joannes", "stanisław", "stanislaus", "wojciech", "adalbertus",
             "woyciech", "jakub", "jacobus", "jacobi", "franciszek", "franciscus",
             "józef", "jozephus", "kazimierz", "casimirus", "marcin", "martinus",
             "andrzej", "andreas", "tomasz", "thomas", "michał", "michael",
             "paweł", "paulus", "piotr", "petrus", "maciej", "mathias"}

def rodzaj_osoby(nazwa):
    """Zwraca 'f' (żeński) lub 'm' (męski) na podstawie nazwy osoby."""
    n = (nazwa or "").strip()
    if not n:
        return "m"
    tokens = n.lower().split()
    # 1) jawne imię żeńskie / męskie
    for t in tokens:
        if t in _IMIONA_F:
            return "f"
        if t in _IMIONA_M:
            return "m"
    last = tokens[-1]
    # 2) żeńskie sufiksy nazwisk/przezwisk mężatek i panien
    if re.search(r"(owa|yna|ina|anka|ówna|owna|cczanka|szczanka|czanka)$", last):
        return "f"
    if re.search(r"(ska|cka|dzka|żanka|zanka)$", last):
        return "f"
    # 3) ogólna końcówka -a zwykle żeńska (np. Zerkowska, Dziedzicowicowa)
    if last.endswith("a") and not last.endswith("ski") and len(last) > 3:
        return "f"
    return "m"

def rola_strona(nazwa, typ):
    """typ: 'wnosi' (powód) lub 'przeciw' (pozwany). Rzeczownik odmieniony rodzajowo."""
    f = rodzaj_osoby(nazwa) == "f"
    if typ == "wnosi":
        return "powódka" if f else "powód"
    return "pozwana" if f else "pozwany"

def redaguj_notatke(src):
    temat = src.get("temat", "").strip()
    kto = src.get("kto", "").strip()
    zkim = src.get("zkim", "").strip()
    okol = src.get("okol", "").strip()
    ludzie = src.get("ludzie", [])

    # --- STRONY: zdanie z deklinacją rodzajową (charakter jest już w nagłówku karty) ---
    strony = ""
    if kto and zkim:
        strony = (f"{kto} ({rola_strona(kto,'wnosi')}) "
                  f"oraz {zkim} ({rola_strona(zkim,'przeciw')}).")
    elif kto:
        strony = f"{kto} ({rola_strona(kto,'wnosi')})."
    elif zkim:
        strony = f"{zkim} ({rola_strona(zkim,'przeciw')})."

    # --- OKOLICZNOŚCI: wszystkie „Dodatkowe okoliczności” z bloku (pierwszy i kolejne wiersze) ---
    # Każdy wpis może mieć kilka okoliczności dopisanych przy różnych osobach — sklejamy je
    # w jeden opis, przypisując podmiot tam, gdzie zdanie zaczyna się od czasownika/zaimka.
    powod_nazwa = kto or ""
    def biernik_f(nazwa):
        n = nazwa.strip()
        if n.endswith("owa"):
            return n[:-3] + "ową"
        if n.endswith("ska") or n.endswith("cka") or n.endswith("dzka"):
            return n[:-1] + "ą"
        if n.endswith("a"):
            return n[:-1] + "ą"
        return n
    VERB_INIT = re.compile(
        r"^(nie było go|nie było jej|był |była |grał|grała|wstąpił|wstąpiła|"
        r"stanął|stanęła|bronił|broniła|szukał|szukała|przyszedł|przyszła|"
        r"został|została|grali|mieszka|mieszkał|poszedł|poszła|uderzył|uderzyła|"
        r"zeznał|zeznała|oświadczył|oświadczyła|widział|widziała|usłyszał)",
        re.IGNORECASE)
    zdania = []
    okol_list = src.get("okol_list", [])
    if not okol_list and okol:
        okol_list = [("", okol)]  # fallback, gdyby lista nie została zebrana
    seen = set()
    for (osoba_w, tekst) in okol_list:
        t = (tekst or "").strip()
        if not t or t in seen:
            continue
        seen.add(t)
        # zamień zaimek „ją/jej” na nazwisko powódki (biernik/dopełniacz), gdy jednoznaczne
        if powod_nazwa and rodzaj_osoby(powod_nazwa) == "f":
            t = re.sub(r"\bją\b", biernik_f(powod_nazwa), t, count=1)
            t = re.sub(r"\bjej\b", powod_nazwa, t, count=1)
        # jeśli okoliczność dopisana przy konkretnej osobie zaczyna się od czasownika,
        # poprzedź ją tą osobą jako podmiotem (np. mąż: „Nie było go…”)
        osoba_w = (osoba_w or "").strip()
        if osoba_w and osoba_w not in (kto, zkim) and VERB_INIT.match(t):
            t = f"{osoba_w}: {t[0].lower()}{t[1:]}"
        zdania.append(zdanie(t))
    okol_txt = " ".join(z for z in zdania if z).strip()

    return {"strony": strony, "okol": okol_txt}

def pnorm(s):
    s = (s or "").lower()
    s = s.replace("ł", "l")
    import unicodedata
    s = "".join(c for c in unicodedata.normalize("NFD", s)
                if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9 ]", " ", s).strip()

RODZINA_RE = re.compile(
    r"\b(brat|brata|bracia|siostr|ojc|ojciec|matk|matka|syn|syna|córk|cork|"
    r"żon|zon|mąż|maz|małżon|malzon|teść|tesc|teści|tesci|szwagr|zięć|ziec|"
    r"wdow|pasierb|macoch|ojczym|wnuk|dziadek|babk|krewn|rodzic|rodzeństw|rodzenstw)",
    re.IGNORECASE)

def is_junk(n):
    n = (n or "").strip()
    return (not n) or n in ("?", "-", "—") or len(re.sub(r"[^0-9A-Za-zĄąĆćĘꣳŃńÓóŚśŻżŹź]", "", n)) < 3

# ---------- role ----------
# Jawna kolumna Rola ma pierwszeństwo; bez niej rolę wyznacza Kto / Z kim.
ROLA_JAWNA = [("pokrzyw", "powód / pokrzywdzony"), ("oskarżyc", "powód / oskarżyciel"),
              ("oskarż", "pozwany / oskarżony"), ("apelant", "apelant"), ("świad", "świadek"),
              ("elekt", "elekt"), ("pełnomoc", "pełnomocnik")]
def rola_osoby(rola, osoba, kto, zkim):
    rl = (rola or "").strip().lower()
    for pref, nazwa in ROLA_JAWNA:
        if rl.startswith(pref):
            return nazwa
    if strona_sprawy(osoba, kto):
        return "powód / wnoszący"
    if strona_sprawy(osoba, zkim):
        return "pozwany / strona przeciwna"
    return "osoba wzmiankowana"
def strona_sprawy(osoba, pole):
    """Czy osoba jest wymieniona w polu Kto / Z kim (cała nazwa albo samo nazwisko z pola)."""
    po = pnorm(osoba)
    if len(po) < 3:
        return False
    for czesc in re.split(r"[,;]|\s+i\s+|\s+oraz\s+", pole or ""):
        pc = pnorm(czesc)
        if len(pc) >= 3 and (pc == po or po.endswith(" " + pc)):
            return True
    return False
def _rola_rank(rola):
    r = (rola or "").lower()
    if "powód" in r or "wnosz" in r or "pokrzyw" in r: return 0
    if "pozwan" in r or "przeciw" in r or "oskarż" in r: return 1
    if "apelant" in r: return 2
    if "świad" in r: return 3
    return 4

# ---------- 1. odczyt arkusza ----------
wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
it = ws.iter_rows(values_only=True)
hdr = [("" if h is None else str(h).strip()) for h in next(it)]
C = {}
for k in KOLUMNY:
    if k not in hdr:
        sys.exit(f"Brak kolumny „{k}” w arkuszu {SRC}")
    C[k] = hdr.index(k)
def v(r, k):
    i = C[k]
    return "" if i >= len(r) or r[i] is None else str(r[i]).strip()
TAK = lambda s: (s or "").strip().lower() in ("tak", "true", "1", "x")

bloki, cur = [], None
for r in it:
    if v(r, "ID"):
        try:
            data = f"{int(v(r,'Rok')):04d}-{int(v(r,'Miesiąc')):02d}-{int(v(r,'Dzień')):02d}"
        except ValueError:
            data = v(r, "Rok")
        cur = {"data": data, "skan": v(r, "Skan"), "strona": v(r, "Strona"), "kto": v(r, "Kto"),
               "zkim": v(r, "Z kim"), "temat": v(r, "Temat sprawy"), "charakter": v(r, "Charakter wpisu"),
               "slowa": v(r, "Słowa klucze"), "okol": v(r, "Dodatkowe okoliczności"),
               "mieso": TAK(v(r, "Mięso?")), "kobieta": TAK(v(r, "Kobieta sprawcza?")),
               "okol_list": [], "ludzie": []}
        bloki.append(cur)
    if cur is None:
        continue
    osoba = v(r, "Osoby powiązane")
    if v(r, "Dodatkowe okoliczności"):
        cur["okol_list"].append((osoba, v(r, "Dodatkowe okoliczności")))
    if not v(r, "ID"):
        cur["mieso"] = cur["mieso"] or TAK(v(r, "Mięso?"))
        cur["kobieta"] = cur["kobieta"] or TAK(v(r, "Kobieta sprawcza?"))
    if any(v(r, k) for k in ("Osoby powiązane", "Stan społeczny", "Zawód", "Uwagi", "Skąd?", "Położenie domu")):
        cur["ludzie"].append({"osoba": osoba, "stan": v(r, "Stan społeczny"), "zaw": v(r, "Zawód"),
                              "uw": v(r, "Uwagi"), "rola": v(r, "Rola"), "skad": v(r, "Skąd?"),
                              "wiek": v(r, "Wiek"), "martwy": TAK(v(r, "Martwy?")), "dom": v(r, "Położenie domu")})
wb.close()

# Strony z kolumn Kto / Z kim, których nie ma wśród osób bloku (np. cech, magistrat
# albo strona wpisana tylko w nagłówku), dopisujemy jako osoby, żeby nie zniknęły z wpisu.
dopisane_strony = 0
for e in bloki:
    for pole in ("kto", "zkim"):
        for czesc in re.split(r"[,;]|\s+i\s+|\s+oraz\s+", e[pole]):
            czesc = czesc.strip()
            if is_junk(czesc):
                continue
            if not any(pnorm(l["osoba"]) == pnorm(czesc) or pnorm(l["osoba"]).endswith(" " + pnorm(czesc))
                       for l in e["ludzie"] if l["osoba"]):
                e["ludzie"].append({"osoba": czesc, "stan": "", "zaw": "", "uw": "", "rola": "", "skad": "",
                                    "wiek": "", "martwy": False, "dom": ""})
                dopisane_strony += 1

# ---------- 2. stałe identyfikatory ----------
poprz = defaultdict(list)
if os.path.exists(OUT):
    for w in json.load(open(OUT, encoding="utf-8")).get("wpisy", []):
        poprz[(w["data"], str(w["skan"]))].append(w["id"])
uzyte, zajete = Counter(), {i for l in poprz.values() for i in l}
def nowe_id(data):
    n = 1
    while f"{data}_{n:02d}" in zajete:
        n += 1
    zajete.add(f"{data}_{n:02d}")
    return f"{data}_{n:02d}"

red = json.load(open(RED, encoding="utf-8")) if os.path.exists(RED) else {}

# ---------- 3. wpisy ----------
wpisy, nowe, red_uzyte = [], 0, 0
for e in bloki:
    k = (e["data"], e["skan"]); i = uzyte[k]; uzyte[k] += 1
    if i < len(poprz[k]):
        wid = poprz[k][i]
    else:
        wid = nowe_id(e["data"]); nowe += 1
    nota = redaguj_notatke(e)
    okol_key = re.sub(r"\s+", " ", " ".join(t for (_o, t) in e["okol_list"]).lower()).strip()
    if okol_key and okol_key in red:
        nota["okol"] = red[okol_key]; red_uzyte += 1
    osoby = []
    for l in e["ludzie"]:
        nazwa = l["osoba"]
        nn = is_junk(nazwa)
        if nn and not (l["stan"] or l["zaw"] or l["uw"] or l["dom"]):
            continue
        if nn:
            imie, nazwisko = "", "NN"
        else:
            parts = nazwa.split()
            imie = parts[0] if len(parts) > 1 else ""
            nazwisko = " ".join(parts[1:]) if len(parts) > 1 else nazwa
        uwagi = " · ".join(x for x in (
            l["zaw"], ("ze " + l["skad"]) if l["skad"] else "", l["uw"],
            l["wiek"].lower() if l["wiek"] else "", "zmarły" if l["martwy"] else "",
            ("dom: " + l["dom"][0].lower() + l["dom"][1:]) if l["dom"] else "") if x)
        osoby.append({"imie": imie, "nazwisko": nazwisko,
                      "rola": rola_osoby(l["rola"], nazwa, e["kto"], e["zkim"]),
                      "status": l["stan"], "miejscowosc": l["skad"], "uwagi": uwagi,
                      **({"nn": True} if nn else {})})
    osoby.sort(key=lambda o: _rola_rank(o["rola"]))
    wpisy.append({
        "id": wid, "skan": e["skan"], "strona": e["strona"], "data": e["data"],
        "typ": e["charakter"], "streszczenie": "", "pewnosc": "wysoka",
        "zrodlo": "radziecka 1712–1713" if e["data"][:4] < "1714" else "radziecka 1724–1736",
        "temat": e["temat"], "charakter": e["charakter"], "slowa": ", ".join(split_kw(e["slowa"])),
        "okol": e["okol"], "nota_strony": nota["strony"], "nota_okol": nota["okol"],
        "kw": split_kw(e["slowa"]), "osoby": osoby,
        **({"mieso": True} if e["mieso"] else {}), **({"kobieta_sprawcza": True} if e["kobieta"] else {}),
        "_ludzie": e["ludzie"],
    })
wpisy.sort(key=lambda w: (w["data"], w["id"]))

# ---------- 4. indeks osób i miejsc ----------
P = {}
miejsca = defaultdict(lambda: {"strony": set(), "liczba": 0})
zawody_licz = Counter()
for w in wpisy:
    etyk = (w["temat"] or w["charakter"] or "sprawa") + (f" ({w['data']})" if w["data"] else "")
    nazwy = [l["osoba"] for l in w["_ludzie"] if not is_junk(l["osoba"])]
    for l in w["_ludzie"]:
        nazwa = l["osoba"]
        if is_junk(nazwa):
            continue
        rec = P.setdefault(nazwa, {"osoba": nazwa, "wpisy": [], "strony": [], "statusy": set(), "zawody": set(),
                                   "uwagi": [], "rodzina": [], "okolicznosci": [],
                                   "pow": defaultdict(lambda: {"ile": 0, "sprawy": []})})
        if w["id"] not in rec["wpisy"]:
            rec["wpisy"].append(w["id"])
        if w["strona"] and w["strona"] not in rec["strony"]:
            rec["strony"].append(w["strona"])
        if l["stan"]: rec["statusy"].add(l["stan"])
        if l["zaw"]:
            rec["zawody"].add(l["zaw"]); zawody_licz[l["zaw"]] += 1
        for tekst in (l["uw"], ("dom: " + l["dom"]) if l["dom"] else ""):
            if tekst:
                lista = rec["rodzina"] if RODZINA_RE.search(tekst) and not tekst.startswith("dom:") else rec["uwagi"]
                if tekst not in lista: lista.append(tekst)
        if l["skad"]:
            n = f"pochodzenie: {l['skad']}"
            if n not in rec["uwagi"]: rec["uwagi"].append(n)
            miejsca[l["skad"]]["strony"].add(w["strona"]); miejsca[l["skad"]]["liczba"] += 1
        if w["okol"]:
            n = f"{etyk}: {w['okol']}"
            if n not in rec["okolicznosci"]: rec["okolicznosci"].append(n)
        for inna in nazwy:
            if inna == nazwa:
                continue
            d = rec["pow"][inna]
            if not any(s["id"] == w["id"] for s in d["sprawy"]):
                d["ile"] += 1; d["sprawy"].append({"label": etyk, "id": w["id"]})
klucz_str = lambda s: (len(s), s)
osoby_indeks = []
for rec in P.values():
    n = len(rec["wpisy"])
    osoby_indeks.append({
        "osoba": rec["osoba"], "liczba": str(n), "liczba_n": n,
        "strony": ", ".join(sorted(rec["strony"], key=klucz_str)),
        "statusy": ", ".join(sorted(rec["statusy"])), "zawody": sorted(rec["zawody"]),
        "uwagi": rec["uwagi"][:12], "rodzina": rec["rodzina"][:12], "okolicznosci": rec["okolicznosci"][:12],
        "powiazania": [{"osoba": k, "ile": d["ile"], "sprawy": d["sprawy"][:5]}
                       for k, d in sorted(rec["pow"].items(), key=lambda kv: -kv[1]["ile"])[:25]],
        "_zaw_norm": pnorm(" ".join(sorted(rec["zawody"]))),
    })
osoby_indeks.sort(key=lambda o: (-o["liczba_n"], o["osoba"].lower()))
miejsca_indeks = [{"miejsce": m, "liczba": str(d["liczba"]),
                   "strony": ", ".join(sorted({s for s in d["strony"] if s}, key=klucz_str))}
                  for m, d in sorted(miejsca.items(), key=lambda kv: kv[0].lower())]

# ---------- 5. ujednolicenie diakrytyków nazwisk w notatkach ----------
def _strip_diac(s):
    s = s.replace("ł", "l").replace("Ł", "L")
    return "".join(c for c in unicodedata.normalize("NFD", s)
                   if unicodedata.category(c) != "Mn")
_diac_forms = defaultdict(lambda: defaultdict(int))
def _collect_name(s):
    for tok in re.split(r"[\s,.;:()„”\"'/–—-]+", s or ""):
        tok = tok.strip()
        if tok and re.search(r"[A-Za-zĄąĆćĘꣳŃńÓóŚśŻżŹź]", tok):
            _diac_forms[_strip_diac(tok).lower()][tok] += 1
for w in wpisy:
    for o in w["osoby"]:
        _collect_name(o["imie"]); _collect_name(o["nazwisko"])
_diac_map = {}
for _k, _forms in _diac_forms.items():
    _withd = {f: n for f, n in _forms.items() if _strip_diac(f) != f}
    if _withd:
        _diac_map[_k] = max(_withd.items(), key=lambda kv: kv[1])[0]
# rdzeń + litera przed 'ski'. Budujemy: rdzen_ascii -> rdzen_z_diakrytykiem.
_FLEKS = ("ego", "emu", "imi", "ich", "iej", "im", "ym", "ej", "ą", "a", "e", "i", "y", "ów", "om", "ami")
_rdzen_map = {}
for _k, _form in _diac_map.items():
    # interesują nas nazwiska na -ski/-cki/-ska/-cka (odmienne przymiotnikowo)
    m = re.search(r"(sk|ck)(i|a)$", _k)
    if m:
        rdzen_ascii = _k[:m.start()+2]          # ...sk / ...ck
        rdzen_diac = _form[:len(_form) - (len(_k) - (m.start()+2))]
        if _strip_diac(rdzen_diac).lower() != rdzen_diac.lower():
            _rdzen_map[rdzen_ascii] = rdzen_diac

def _popraw_token(tok):
    """Zwraca poprawioną pisownię tokenu lub None, gdy nic nie zmieniać."""
    if _strip_diac(tok) != tok:            # już ma diakrytyk
        return None
    low = tok.lower()
    # 1) pełne dopasowanie (mianownik)
    corr = _diac_map.get(low)
    if corr and corr.lower() != low:
        return corr[0].upper() + corr[1:] if tok[0].isupper() else corr
    # 2) forma odmieniona nazwiska na -sk-/-ck- : dopasuj rdzeń, zachowaj końcówkę
    m = re.search(r"(sk|ck)", low)
    if m:
        for end in range(len(low), m.end()-1, -1):
            rdzen = low[:end]
            if rdzen in _rdzen_map and rdzen != _rdzen_map[rdzen].lower():
                nowy = _rdzen_map[rdzen] + tok[end:]   # rdzeń z diakrytykiem + reszta
                return nowy[0].upper() + nowy[1:] if tok[0].isupper() else nowy
    return None

def norm_diakrytyki(txt):
    def _repl(m):
        tok = m.group(0)
        nowy = _popraw_token(tok)
        return nowy if nowy is not None else tok
    return re.sub(r"[A-Za-zĄąĆćĘꣳŃńÓóŚśŻżŹź]+", _repl, txt or "")

for w in wpisy:
    w["nota_okol"] = norm_diakrytyki(w["nota_okol"])
    w["nota_strony"] = norm_diakrytyki(w["nota_strony"])
    w.pop("_ludzie", None)

# ---------- 6. zapis ----------
kw_licz = Counter(k for w in wpisy for k in w["kw"])
slowa_klucze = sorted(kw_licz, key=lambda t: (-kw_licz[t], t.lower()))
licz = lambda pole: Counter(w[pole] for w in wpisy if w[pole] and w[pole] != "-")
char_l, tem_l = licz("charakter"), licz("temat")
zawody = sorted(zawody_licz, key=lambda z: (-zawody_licz[z], z.lower()))
out = {
    "zrodlo": "Księga radziecka Wiśnicza 1712–1736",
    "wpisy": wpisy, "osoby_indeks": osoby_indeks, "miejsca_indeks": miejsca_indeks,
    "meta": {
        "lata": sorted({w["data"][:4] for w in wpisy if w["data"][:4].isdigit()}),
        "typy": sorted({w["typ"] for w in wpisy if w["typ"]}),
        "charaktery": sorted(char_l, key=lambda c: (-char_l[c], c.lower())),
        "tematy": sorted(tem_l, key=lambda t: (-tem_l[t], t.lower())),
        "role": sorted({o["rola"] for w in wpisy for o in w["osoby"]}, key=lambda r: (_rola_rank(r), r)),
        "slowa_klucze": slowa_klucze, "slowa_klucze_top": slowa_klucze[:60],
        "zawody": zawody, "zawody_top": zawody[:60],
        "liczba_wpisow": len(wpisy), "liczba_osob_wierszy": sum(len(w["osoby"]) for w in wpisy),
        "liczba_osob_indeks": len(osoby_indeks), "liczba_miejsc": len(miejsca_indeks),
    },
}
os.makedirs(os.path.dirname(OUT), exist_ok=True)
with open(OUT, "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, separators=(",", ":"))
print(f"Wpisów: {len(wpisy)} (nowych: {nowe}), osób w indeksie: {len(osoby_indeks)}, "
      f"miejsc: {len(miejsca_indeks)}, haseł: {len(slowa_klucze)}, zredagowanych notatek: {red_uzyte}, "
      f"stron dopisanych z Kto / Z kim: {dopisane_strony}")
